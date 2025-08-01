from io import BytesIO
from typing import Tuple
from openpyxl import load_workbook
from ..constants import NAME_TOKEN

class ExcelService:
    @staticmethod
    def anonymize_name(file_obj, sheet_name: str | None = None) -> Tuple[str, BytesIO]:
        """
        Substitui **apenas** valores da coluna 'nome' (case-insensitive) por `NAME_TOKEN`.
        Mantém layout original.
        """
        wb = load_workbook(file_obj)
        ws = wb[sheet_name] if sheet_name else wb.active

        # Localiza coluna Nome
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=False))
        name_col_idx = None
        for cell in header_row:
            if str(cell.value).strip().lower() == "nome":
                name_col_idx = cell.column
                break
        if name_col_idx is None:
            raise ValueError("Coluna 'Nome' não encontrada.")

        # Anonimiza células
        for cell in ws.iter_cols(min_col=name_col_idx, max_col=name_col_idx,
                                 min_row=2, values_only=False)[0]:
            if cell.value:
                cell.value = NAME_TOKEN

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f"anon_nome_{file_obj.name}"
        return fname, buf
