from io import BytesIO
from typing import Dict, List, Tuple, TypedDict

from openpyxl import load_workbook

class SheetPreview(TypedDict):
    name: str
    columns: List[str]

class PreviewResponse(TypedDict):
    sheets: List[SheetPreview]
    tokens_detected: List[str]
from openpyxl.worksheet.worksheet import Worksheet

from ..tokens import TOKENS, discover_tokens_in_headers


class AnonimizationService:
    @staticmethod
    def scan_headers(ws: Worksheet) -> List[str]:
        return [str(c.value).strip() if c.value else "" for c in ws[1]]

    @classmethod
    def preview_anonymization(cls, file_obj) -> PreviewResponse:
        wb = load_workbook(file_obj, read_only=True, data_only=False)
        sheets_preview: List[SheetPreview] = []
        tokens_detected: set[str] = set()

        for ws in wb.worksheets:
            headers = cls.scan_headers(ws)
            tokens_in_sheet = discover_tokens_in_headers(headers)
            sheets_preview.append({"name": ws.title, "columns": headers})
            tokens_detected.update(tokens_in_sheet.keys())

        return {
            "sheets": sheets_preview,
            "tokens_detected": sorted(list(tokens_detected)),
        }

    @classmethod
    def run_anonymization(cls, file_obj, rules: Dict[str, bool]) -> Tuple[BytesIO, str]:
        wb = load_workbook(file_obj, data_only=False)
        for ws in wb.worksheets:
            headers = cls.scan_headers(ws)
            tokens_to_apply = discover_tokens_in_headers(headers)

            for token_key, col_indices in tokens_to_apply.items():
                if not rules.get(token_key, False):
                    continue  # Skip if rule is false or absent

                _, mask_fn = TOKENS[token_key]
                for row in ws.iter_rows(min_row=2):
                    for col_idx in col_indices:
                        cell = row[col_idx]
                        if cell.value is not None:
                            cell.value = mask_fn(str(cell.value))

        out = BytesIO()
        wb.save(out)
        out.seek(0)
        return out, f"anonimized_{getattr(file_obj, 'name', 'file.xlsx')}"
