import json
from io import BytesIO
from typing import Dict

from django.http import HttpResponse, JsonResponse
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt
from openpyxl import load_workbook
from rest_framework.parsers import MultiPartParser, JSONParser
from rest_framework.views import APIView

from .tokens import TOKENS, discover_tokens_in_headers


def scan_headers(ws) -> list[str]:
    return [str(c.value).strip() if c.value else "" for c in ws[1]]


@method_decorator(csrf_exempt, name="dispatch")
class AnonymizePreview(APIView):
    parser_classes = (MultiPartParser,)

    def post(self, request, *a, **kw):
        f = request.data.get("file")
        if not f:
            return HttpResponse("arquivo não enviado", status=400)

        wb = load_workbook(f, read_only=True, data_only=False)
        resp = {"sheets": [], "tokens_detected": set()}

        for ws in wb.worksheets:
            headers = scan_headers(ws)
            tokens_in_sheet = discover_tokens_in_headers(headers)
            resp["sheets"].append({"name": ws.title, "columns": headers})
            resp["tokens_detected"].update(tokens_in_sheet.keys())

        resp["tokens_detected"] = sorted(list(resp["tokens_detected"]))
        return JsonResponse(resp)


@method_decorator(csrf_exempt, name="dispatch")
class AnonymizeRun(APIView):
    parser_classes = (MultiPartParser, JSONParser)

    def post(self, request, *a, **kw):
        f = request.data.get("file")
        if not f:
            return HttpResponse("arquivo não enviado", status=400)

        raw_rules = request.data.get("rules", "{}")
        rules: Dict[str, bool] = json.loads(raw_rules)
        if not rules:
            return HttpResponse("regras de anonimização não enviadas", status=400)

        wb = load_workbook(f, data_only=False)
        for ws in wb.worksheets:
            headers = scan_headers(ws)
            tokens_to_apply = discover_tokens_in_headers(headers)

            for token_key, col_indices in tokens_to_apply.items():
                if not rules.get(token_key, False):
                    continue  # Skip if rule is false or absent

                _, mask_fn = TOKENS[token_key]
                for row in ws.iter_rows(min_row=2):
                    for col_idx in col_indices:
                        cell = row[col_idx]
                        if cell.value is not None:
                            cell.value = mask_fn(str(cell.value))

        out = BytesIO()
        wb.save(out)
        out.seek(0)
        return HttpResponse(
            out.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename=anonimized_{getattr(f, "name", "file.xlsx")}'
            },
        )
